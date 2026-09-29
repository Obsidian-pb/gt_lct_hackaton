"""Read the supplied workbook; preserve source rows and conditional service cells."""
import argparse
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SERVICE_COLUMNS = {
 'N':'101','O':'101','P':'PSC','Q':'PSC','R':'PSC','S':'PSC','T':'MGPSS',
 'U':'102','V':'102','W':'102','X':'103','Y':'103','Z':'103','AA':'104','AB':'104',
 'AC':'ZEMP','AD':'ZEMP','AE':'ZEMP','AF':'ZEMP','AG':'ZEMP','AH':'FSB','AI':'FSB',
 'AJ':'MOSOBLGAZ','AK':'AUTOROADS','AL':'MOSGORTRANS','AM':'MOSGORTRANS','AN':'MOSGORTRANS',
 'AO':'GKH','AP':'GORMOST','AQ':'GORMOST','AR':'GORMOST','AS':'GORMOST','AT':'CANAL',
 'AU':'MGTS','AV':'MGTS','AW':'METRO','AX':'MOSVODOCANAL','AY':'MOEK','AZ':'MOESK',
 'BA':'OEK','BB':'MOSLIFT','BC':'ZODD','BD':'DEP_GKH','BE':'MOSBEZ','BF':'MOSBEZ_ANALYTICS',
 'BG':'MAYOR','BH':'MOSCOLLECTOR','BI':'MZD','BJ':'EDUCATION','BK':'WATER_REGION',
 'BL':'MILITARY','BM':'OATI','BN':'MOSVODOSTOK','BO':'DepEco','BP':'Dep.tszn',
 'BQ':'RSVO','BR':'EVAZHD','BS':'MSPPN','BT':'RITUAL','BU':'DTU','BV':'ROSGVARDIA',
 'BW':'OIV','BX':'OIV_TINAO','BY':'AUTOROADS_AO','BZ':'CONSTRUCTION','CA':'CONSTRUCTION',
 'CB':'VETERINARY','CC':'MOSZHIL','CD':'CULTURE','CE':'GLINKA','CF':'NTU','CG':'FSO',
 'CH':'MSR','CI':'MSR','CJ':'TOURISM','CK':'URBAN','CL':'URBAN','CM':'MOD',
 'CN':'MOD_UAV','CO':'TRANSPORT','CP':'TRANSPORT','CQ':'ECO_MONITOR','CR':'MOD_RHBZ','CS':'MOD_RHBZ',
 'CT':'CITYENERGY','CU':'CIVIL_CONSTRUCTION'}
MAIN_ALIASES={'MCHS':'101','Police':'102','AMBULANCE':'103','MOSGAZ':'104','МСР':'MSR'}
def clean(v):
    return str(v).strip() if v is not None else ''

def convert(path):
    ws=load_workbook(path,data_only=True).active
    def header(row,col):
        value=ws.cell(row,col).value
        if value is None:
            for merged in ws.merged_cells.ranges:
                if merged.min_row<=row<=merged.max_row and merged.min_col<=col<=merged.max_col:
                    value=ws.cell(merged.min_row,merged.min_col).value;break
        return clean(value)
    groups={};current='';services={};columns=[];entries=[]
    for col in range(14,ws.max_column+1):
        letter=get_column_letter(col);sid=SERVICE_COLUMNS[letter]
        name=header(1,col)
        if sid in ('101','PSC','MGPSS'):name={'101':'Служба 101','PSC':'ОДС ПСЦ','MGPSS':'МГПСС'}[sid]
        if sid in ('102','103','104'):name={'102':'Служба 102 — МВД','103':'Служба 103 — СМП','104':'Служба 104 — МОСГАЗ'}[sid]
        services.setdefault(sid,name)
        condition=' / '.join(dict.fromkeys(x for x in [header(2,col),header(3,col)] if x and x!=header(1,col)))
        columns.append({'column':letter,'service':sid,'condition':condition})
    for r in range(4,ws.max_row+1):
        values=[ws.cell(r,c).value for c in range(1,14)]
        if values[0] is None:
            if values[5] is not None:current=clean(values[5])
            continue
        group_id=str(values[0]);groups.setdefault(group_id,current)
        code=clean(values[4]);main_raw=clean(values[12])
        mains=[MAIN_ALIASES.get(x.strip(),x.strip()) for x in main_raw.split(',') if x.strip()]
        unknown=[x for x in mains if x not in services]
        if unknown:raise ValueError(f'Unknown main service at row {r}: {unknown}')
        rules=[{**item,'value':clean(ws.cell(r,col).value)} for col,item in enumerate(columns,14) if ws.cell(r,col).value is not None]
        entries.append({'id':code,'category':group_id,'group':groups[group_id],
                        'statistical_group':clean(values[5]),'sign1':clean(values[6]),'sign2':clean(values[7]),'sign3':clean(values[8]),
                        'extra_signs':clean(values[9]),'title':clean(values[10]),'ekp_type':clean(values[11]),
                        'main_service':mains[0] if len(mains)==1 else '', 'services':mains,
                        'main_service_source':main_raw,'rules':rules,'source_row':r})
    assert len(entries)==len({x['id'] for x in entries})
    return {'version':'112-v046-24','source':{'file':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheet':ws.title},
            'categories':groups,'services':services,'columns':columns,'entries':entries}

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('workbook',type=Path);args=parser.parse_args()
    data=convert(args.workbook);out=ROOT/'catalog'/'classifier.json';out.parent.mkdir(exist_ok=True)
    out.write_text(json.dumps(data,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
    print(f"Imported {len(data['categories'])} groups, {len(data['entries'])} records, {len(data['services'])} services")
