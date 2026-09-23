from __future__ import annotations

import argparse
import asyncio
import os
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable
from uuid import NAMESPACE_URL, UUID, uuid4, uuid5

import openpyxl
from openpyxl.utils import get_column_letter
from sqlalchemy import delete, func, update
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncConnection, create_async_engine


DATABASE_ROOT = Path(__file__).resolve().parents[1]
REPOSITORY_ROOT = DATABASE_ROOT.parent
sys.path.insert(0, str(DATABASE_ROOT / "src"))

from trainer_db.models.catalog import (  # noqa: E402
    ClassifierVersion,
    EventClass,
    EventServiceRoute,
    EventFeature1,
    EventFeature2,
    EventFeature3,
    EventType,
    Service,
    classifier_version_events,
    event_class_services,
)


DEFAULT_WORKBOOK = (
    REPOSITORY_ROOT
    / "TZ"
    / "datasets"
    / "Классификатор_происшествий_v_046_11_ДТУ_15_11_2024_искл_пожар_задымление.xlsx"
)
EXCLUSION_TEXT = "не отображается оператору 112"
EXPECTED_IMPORTED_ROWS = 1137
EXPECTED_EXCLUDED_ROWS = 144

SERVICE_GROUP_CODES = {
    "Классификатор МЧС": "MCHS",
    "Классификатор МВД": "POLICE",
    "Классификатор СМП": "AMBULANCE",
    "Классификатор МОСГАЗ": "MOSGAZ",
    "ЦЭМП": "ZEMP",
    "Классификатор ФСБ": "FSB",
    "Классификатор Мособлгаз": "MOSOBLGAZ",
    "Автомобильные дороги": "AUTOROADS",
    "Мосгортранс": "MOSGORTRANS",
    "Гор. Хозяйство": "GKH",
    "ГОРМОСТ": "GORMOST",
    "Канал имени Москвы": "MOSCOW_CANAL",
    "МГТС": "MGTS",
    "Метро": "METRO",
    "Мосводоканал": "MOSVODOCANAL",
    "МОЭК": "MOEK",
    'МОЭСК (ПАО "Россети Московский регион")': "MOESK",
    "ОЭК": "OEK",
    "Мослифт": "MOSLIFT",
    "ЦОДД": "ZODD",
    "Деп. ЖКХ": "DEP_GKH",
    "Департамент РБиПК (ГКУ МОСБЕЗ)": "DEP_RBPK",
    "Аппарат МЭРА": "MAYOR_OFFICE",
    "Москоллектор": "MOSCOLLECTOR",
    "РЖД": "MZD",
    "Департамент образования": "DEP_EDUCATION",
    "Центррегионводхоз (Московско-Окское БВУ)": "WATER_AUTHORITY",
    "Военная комендатура": "MILITARY_COMMAND",
    "ОАТИ": "OATI",
    "Мосводосток": "MOSVODOSTOK",
    "Департамент ППиООС": "DEP_ECO",
    "ОД Департамент ТСЗН": "DEP_TSZN",
    "РСВО": "RSVO",
    "ЭВАЖД": "EVAZHD",
    "МСППН": "MSPPN",
    "ДТУ_Р (Ритуал)": "DTU_R",
    "ДТУ": "DTU",
    "Росгвардия": "ROSGVARDIA",
    "Территориальные ОИВ": "TERRITORIAL_OIV",
    "Территориальные ОИВ ТиНАО": "TINAO_OIV",
    "Автомобильные дороги АО г.Москвы": "AO_AUTOROADS",
    "Департамент строительства города Москвы": "DEP_CONSTRUCTION",
    "Комитет ветеринарии": "VETERINARY_COMMITTEE",
    "Мосжилинспекция": "MOSZHILINSPEKTSIYA",
    "Департамент культуры": "DEP_CULTURE",
    "ГКУ ЦСА имени Е.П.Глинки": "GLINKA_CENTER",
    "ГКУ НТУ": "GKU_NTU",
    "ФСО": "FSO",
    "ГУП МСР": "GUP_MSR",
    "Комитет по туризму г.Москвы": "TOURISM_COMMITTEE",
    "ДГП (Департамент градостроительной политики)": "DEP_URBAN_POLICY",
}
MAIN_SERVICE_CODES = {
    "Police": "POLICE",
    "Dep.tszn": "DEP_TSZN",
    "DepEco": "DEP_ECO",
}
MAIN_SERVICE_NAMES = {
    "MCHS": "МЧС",
    "POLICE": "Полиция",
    "AMBULANCE": "Скорая медицинская помощь",
    "MOSGAZ": "Мосгаз",
    "ZEMP": "ЦЭМП",
}
CONDITION_BY_COLUMN = {
    "O": "no_access_false", "P": "no_access", "Q": "default",
    "R": "threat_to_people", "S": "victims", "T": "no_access",
    "V": "no_violation_or_victims", "W": "law_violation", "X": "victims",
    "Y": "no_victims", "Z": "victims", "AA": "victims_not_on_scene",
    "AB": "default", "AC": "gasified", "AD": "default",
    "AE": "threat_to_people", "AF": "victims", "AG": "medical_help",
    "AH": "evacuation", "AI": "default", "AJ": "many_people",
    "AM": "default", "AN": "victims", "AO": "traffic_blocked",
    "AQ": "default", "AR": "tunnel", "AS": "pedestrian",
    "AT": "automotive", "AV": "always", "AW": "communication_facility",
    "CA": "default", "CB": "construction", "CE": "listed_cultural_site",
}


@dataclass(frozen=True, slots=True)
class ServiceRoute:
    service_code: str
    source_column: str
    condition_code: str
    condition_label: str | None
    response_label: str | None
    is_primary: bool = False


@dataclass(frozen=True, slots=True)
class ClassifierRow:
    source_row: int
    event_type_code: int
    feature_1_code: int
    feature_2_code: int
    feature_3_code: int
    event_type_name: str
    feature_1_label: str
    feature_2_label: str | None
    feature_3_label: str | None
    main_service_code: str | None
    service_routes: tuple[ServiceRoute, ...]

    @property
    def event_number(self) -> int:
        return (
            self.event_type_code * 1_000_000
            + self.feature_1_code * 10_000
            + self.feature_2_code * 100
            + self.feature_3_code
        )

    @property
    def key(self) -> tuple[int, int, int, int]:
        return (
            self.event_type_code,
            self.feature_1_code,
            self.feature_2_code,
            self.feature_3_code,
        )


@dataclass(frozen=True, slots=True)
class ParsedClassifier:
    rows: tuple[ClassifierRow, ...]
    excluded_rows: int
    group_names: dict[int, str]
    service_names: dict[str, str]


def clean_text(value: object) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).strip().split())
    return text or None


def normalized(value: object) -> str:
    return (clean_text(value) or "").casefold()


def integer_code(value: object, *, cell: str) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{cell}: ожидалось целое число, получено {value!r}")
    result = int(value)
    if result != value:
        raise ValueError(f"{cell}: код должен быть целым, получено {value!r}")
    return result


def service_columns(raw_rows: list[tuple[object, ...]]) -> dict[int, tuple[str, str, str | None]]:
    groups: dict[int, tuple[str, str, str | None]] = {}
    current_group: str | None = None
    for index in range(14, 90):  # O:CL, service classifier columns
        current_group = clean_text(raw_rows[0][index]) or current_group
        if current_group is None or current_group not in SERVICE_GROUP_CODES:
            raise ValueError(f"{get_column_letter(index + 1)}: неизвестная группа службы {current_group!r}")
        condition_label = clean_text(raw_rows[2][index]) or clean_text(raw_rows[1][index])
        groups[index] = (
            SERVICE_GROUP_CODES[current_group], current_group, condition_label
        )
    return groups


def row_service_routes(
    values: tuple[object, ...],
    columns: dict[int, tuple[str, str, str | None]],
    service_names: dict[str, str],
) -> tuple[str | None, tuple[ServiceRoute, ...]]:
    routes: list[ServiceRoute] = []
    main_codes: list[str] = []
    for raw_code in (clean_text(values[13]) or "").split(","):
        raw_code = raw_code.strip()
        if not raw_code:
            continue
        code = MAIN_SERVICE_CODES.get(raw_code, raw_code.upper().replace(".", "_"))
        main_codes.append(code)
        service_names.setdefault(code, MAIN_SERVICE_NAMES.get(code, raw_code))
        routes.append(ServiceRoute(code, "N", "always", "Главная служба", None, len(main_codes) == 1))
    for index, (code, name, condition_label) in columns.items():
        response = clean_text(values[index])
        if response is None or normalized(response) == "нет реагирования":
            continue
        service_names.setdefault(code, MAIN_SERVICE_NAMES.get(code, name))
        column = get_column_letter(index + 1)
        routes.append(ServiceRoute(
            service_code=code,
            source_column=column,
            condition_code=CONDITION_BY_COLUMN.get(column, "always"),
            condition_label=condition_label,
            response_label=response,
        ))
    if not any(
        route.service_code == "AMBULANCE" and route.condition_code == "victims"
        for route in routes
    ):
        service_names.setdefault("AMBULANCE", MAIN_SERVICE_NAMES["AMBULANCE"])
        routes.append(ServiceRoute(
            service_code="AMBULANCE",
            source_column="USR",
            condition_code="victims",
            condition_label="Пострадавшие или погибшие",
            response_label="Вызов скорой медицинской помощи",
        ))
    return (main_codes[0] if main_codes else None), tuple(routes)


def parse_workbook(path: Path) -> ParsedClassifier:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError("Ожидался один лист классификатора")

        worksheet = workbook[workbook.sheetnames[0]]
        raw_rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
        columns = service_columns(raw_rows)
        service_names: dict[str, str] = {}
        group_names: dict[int, str] = {}
        pending_group_name: str | None = None
        prepared: list[tuple[int, tuple[object, ...]]] = []
        excluded_rows = 0

        for row_number, row in enumerate(raw_rows, start=1):
            values = tuple(row)
            first_four = values[:4]
            is_classifier_row = len(first_four) == 4 and all(
                isinstance(value, (int, float)) and not isinstance(value, bool)
                for value in first_four
            )

            if not is_classifier_row:
                if row_number > 3:
                    group_name = clean_text(values[5] if len(values) > 5 else None)
                    group_code = values[4] if len(values) > 4 else None
                    if group_name and isinstance(group_code, (int, float)):
                        group_names[int(group_code)] = group_name
                        pending_group_name = None
                    elif group_name:
                        pending_group_name = group_name
                continue

            event_type_code = integer_code(values[0], cell=f"A{row_number}")
            if event_type_code not in group_names and pending_group_name:
                group_names[event_type_code] = pending_group_name
                pending_group_name = None

            feature_1_label = clean_text(values[6] if len(values) > 6 else None)
            if normalized(feature_1_label) == EXCLUSION_TEXT:
                excluded_rows += 1
                continue
            if not feature_1_label:
                raise ValueError(f"G{row_number}: отсутствует расшифровка признака п1")

            prepared.append((row_number, values))

        missing_group_codes = sorted(
            {
                integer_code(values[0], cell=f"A{row_number}")
                for row_number, values in prepared
            }
            - group_names.keys()
        )
        if missing_group_codes:
            raise ValueError(
                "Не найдены названия групп происшествий для кодов: "
                + ", ".join(map(str, missing_group_codes))
            )

        result_rows: list[ClassifierRow] = []
        seen_numbers: dict[int, int] = {}
        seen_keys: dict[tuple[int, int, int, int], int] = {}

        for row_number, values in prepared:
            codes = (
                integer_code(values[0], cell=f"A{row_number}"),
                integer_code(values[1], cell=f"B{row_number}"),
                integer_code(values[2], cell=f"C{row_number}"),
                integer_code(values[3], cell=f"D{row_number}"),
            )
            main_service_code, routes = row_service_routes(values, columns, service_names)
            if not any(route.source_column != "USR" for route in routes):
                raise ValueError(f"Строка {row_number}: в классификаторе нет ни одной службы")
            item = ClassifierRow(
                source_row=row_number,
                event_type_code=codes[0],
                feature_1_code=codes[1],
                feature_2_code=codes[2],
                feature_3_code=codes[3],
                event_type_name=group_names[codes[0]],
                feature_1_label=clean_text(values[6]) or "",
                feature_2_label=clean_text(values[7]),
                feature_3_label=clean_text(values[8]),
                main_service_code=main_service_code,
                service_routes=routes,
            )
            if item.key in seen_keys:
                raise ValueError(
                    f"Строки {seen_keys[item.key]} и {row_number}: "
                    f"повтор комбинации {item.key}"
                )
            if item.event_number in seen_numbers:
                raise ValueError(
                    f"Строки {seen_numbers[item.event_number]} и {row_number}: "
                    f"повтор номера {item.event_number}"
                )
            seen_keys[item.key] = row_number
            seen_numbers[item.event_number] = row_number
            result_rows.append(item)

        if len(result_rows) != EXPECTED_IMPORTED_ROWS:
            raise ValueError(
                f"После фильтра ожидалось {EXPECTED_IMPORTED_ROWS} строк, "
                f"получено {len(result_rows)}"
            )
        if excluded_rows != EXPECTED_EXCLUDED_ROWS:
            raise ValueError(
                f"Ожидалось исключить {EXPECTED_EXCLUDED_ROWS} строк, "
                f"исключено {excluded_rows}"
            )

        return ParsedClassifier(
            rows=tuple(result_rows),
            excluded_rows=excluded_rows,
            group_names=group_names,
            service_names=service_names,
        )
    finally:
        workbook.close()


def first_labels(
    rows: Iterable[ClassifierRow],
) -> tuple[
    dict[tuple[int, int], str],
    dict[tuple[int, int, int], str | None],
    dict[tuple[int, int, int, int], str | None],
]:
    feature_1: dict[tuple[int, int], str] = {}
    feature_2: dict[tuple[int, int, int], str | None] = {}
    feature_3: dict[tuple[int, int, int, int], str | None] = {}
    for row in rows:
        key_1 = (row.event_type_code, row.feature_1_code)
        key_2 = (*key_1, row.feature_2_code)
        key_3 = (*key_2, row.feature_3_code)
        feature_1.setdefault(key_1, row.feature_1_label)
        if key_2 not in feature_2 or feature_2[key_2] is None:
            feature_2[key_2] = row.feature_2_label
        if key_3 not in feature_3 or feature_3[key_3] is None:
            feature_3[key_3] = row.feature_3_label
    return feature_1, feature_2, feature_3


def sql_text(value: str | None) -> str:
    if value is None:
        return "NULL"
    return "'" + value.replace("'", "''") + "'"


def stable_id(entity: str, *parts: object) -> UUID:
    identity = ":".join(["system112", entity, *(str(part) for part in parts)])
    return uuid5(NAMESPACE_URL, identity)


def write_sql_import(
    output_path: Path,
    parsed: ParsedClassifier,
    workbook_path: Path,
    version_number: int,
) -> None:
    feature_1_labels, feature_2_labels, feature_3_labels = first_labels(parsed.rows)
    version_id = stable_id("classifier-version", version_number)
    active_type_codes = {row.event_type_code for row in parsed.rows}
    lines = [
        "BEGIN;",
        (
            "UPDATE catalog.classifier_versions SET is_active = false, "
            f"updated_at = CURRENT_TIMESTAMP WHERE version_number <> {version_number};"
        ),
        (
            "INSERT INTO catalog.classifier_versions "
            "(id, version_number, name, source_name, description, is_active, published_at) "
            f"VALUES ('{version_id}'::uuid, {version_number}, "
            f"{sql_text('Классификатор происшествий v_046_11')}, "
            f"{sql_text(workbook_path.name)}, "
            f"{sql_text('Импортированы A-D, G-I и маршруты служб N-CL. Строки с пометкой '
                         "'Не отображается оператору 112' исключены.")}, "
            "true, CURRENT_TIMESTAMP) "
            "ON CONFLICT (version_number) DO UPDATE SET "
            "name = EXCLUDED.name, source_name = EXCLUDED.source_name, "
            "description = EXCLUDED.description, is_active = true, "
            "published_at = EXCLUDED.published_at, updated_at = CURRENT_TIMESTAMP;"
        ),
    ]

    for code, name in sorted(parsed.service_names.items()):
        lines.append(
            "INSERT INTO catalog.services (id, code, name) "
            f"VALUES ('{stable_id('service', code)}'::uuid, {sql_text(code)}, {sql_text(name)}) "
            "ON CONFLICT (code) DO NOTHING;"
        )

    for code, name in sorted(parsed.group_names.items()):
        if code not in active_type_codes:
            continue
        entity_id = stable_id("event-type", code)
        lines.append(
            "INSERT INTO catalog.event_types (id, code, name) "
            f"VALUES ('{entity_id}'::uuid, {code}, {sql_text(name)}) "
            "ON CONFLICT (code) DO UPDATE SET name = EXCLUDED.name, "
            "updated_at = CURRENT_TIMESTAMP;"
        )

    for key, label in sorted(feature_1_labels.items()):
        type_code, code = key
        entity_id = stable_id("feature-1", *key)
        lines.append(
            "INSERT INTO catalog.event_features_1 "
            "(id, event_type_id, code, statistics_name, operator_label) "
            f"SELECT '{entity_id}'::uuid, id, {code}, {sql_text(label)}, "
            f"{sql_text(label)} FROM catalog.event_types WHERE code = {type_code} "
            "ON CONFLICT (event_type_id, code) DO UPDATE SET "
            "statistics_name = EXCLUDED.statistics_name, "
            "operator_label = EXCLUDED.operator_label, "
            "updated_at = CURRENT_TIMESTAMP;"
        )

    for key, label in sorted(feature_2_labels.items()):
        type_code, feature_1_code, code = key
        entity_id = stable_id("feature-2", *key)
        lines.append(
            "INSERT INTO catalog.event_features_2 "
            "(id, event_feature_1_id, code, name) "
            f"SELECT '{entity_id}'::uuid, f1.id, {code}, {sql_text(label)} "
            "FROM catalog.event_features_1 AS f1 "
            "JOIN catalog.event_types AS et ON et.id = f1.event_type_id "
            f"WHERE et.code = {type_code} AND f1.code = {feature_1_code} "
            "ON CONFLICT (event_feature_1_id, code) DO UPDATE SET "
            "name = EXCLUDED.name, updated_at = CURRENT_TIMESTAMP;"
        )

    for key, label in sorted(feature_3_labels.items()):
        type_code, feature_1_code, feature_2_code, code = key
        entity_id = stable_id("feature-3", *key)
        lines.append(
            "INSERT INTO catalog.event_features_3 "
            "(id, event_feature_2_id, code, name) "
            f"SELECT '{entity_id}'::uuid, f2.id, {code}, {sql_text(label)} "
            "FROM catalog.event_features_2 AS f2 "
            "JOIN catalog.event_features_1 AS f1 ON f1.id = f2.event_feature_1_id "
            "JOIN catalog.event_types AS et ON et.id = f1.event_type_id "
            f"WHERE et.code = {type_code} AND f1.code = {feature_1_code} "
            f"AND f2.code = {feature_2_code} "
            "ON CONFLICT (event_feature_2_id, code) DO UPDATE SET "
            "name = EXCLUDED.name, updated_at = CURRENT_TIMESTAMP;"
        )

    for row in parsed.rows:
        entity_id = stable_id("event-class", row.event_number)
        labels = [
            label
            for label in (
                row.feature_1_label,
                row.feature_2_label,
                row.feature_3_label,
            )
            if label
        ]
        event_name = " / ".join(labels)[:255]
        lines.append(
            "INSERT INTO catalog.event_classes "
            "(id, event_number, event_type_id, event_feature_1_id, "
            "event_feature_2_id, event_feature_3_id, feature_1_label, "
            "feature_2_label, feature_3_label, name, main_service_id, is_active) "
            f"SELECT '{entity_id}'::uuid, {row.event_number}, et.id, f1.id, "
            f"f2.id, f3.id, {sql_text(row.feature_1_label)}, "
            f"{sql_text(row.feature_2_label)}, {sql_text(row.feature_3_label)}, "
            f"{sql_text(event_name)}, "
            + (
                f"(SELECT id FROM catalog.services WHERE code = {sql_text(row.main_service_code)}), "
                if row.main_service_code else "NULL, "
            )
            + "true "
            "FROM catalog.event_types AS et "
            "JOIN catalog.event_features_1 AS f1 ON f1.event_type_id = et.id "
            "JOIN catalog.event_features_2 AS f2 ON f2.event_feature_1_id = f1.id "
            "JOIN catalog.event_features_3 AS f3 ON f3.event_feature_2_id = f2.id "
            f"WHERE et.code = {row.event_type_code} "
            f"AND f1.code = {row.feature_1_code} "
            f"AND f2.code = {row.feature_2_code} "
            f"AND f3.code = {row.feature_3_code} "
            "ON CONFLICT (event_number) DO UPDATE SET "
            "feature_1_label = EXCLUDED.feature_1_label, "
            "feature_2_label = EXCLUDED.feature_2_label, "
            "feature_3_label = EXCLUDED.feature_3_label, "
            "name = EXCLUDED.name, main_service_id = EXCLUDED.main_service_id, "
            "is_active = true, deleted_at = NULL, "
            "purge_after = NULL, updated_at = CURRENT_TIMESTAMP;"
        )

        for route in row.service_routes:
            lines.append(
                "INSERT INTO catalog.event_class_services (event_class_id, service_id) "
                "SELECT ec.id, s.id FROM catalog.event_classes AS ec "
                "JOIN catalog.services AS s ON s.code = "
                f"{sql_text(route.service_code)} "
                f"WHERE ec.event_number = {row.event_number} ON CONFLICT DO NOTHING;"
            )
            lines.append(
                "INSERT INTO catalog.event_service_routes "
                "(id, event_class_id, service_id, source_column, condition_code, "
                "condition_label, response_label, is_primary) "
                f"SELECT '{stable_id('service-route', row.event_number, route.source_column, route.service_code)}'::uuid, "
                "ec.id, s.id, "
                f"{sql_text(route.source_column)}, {sql_text(route.condition_code)}, "
                f"{sql_text(route.condition_label)}, {sql_text(route.response_label)}, "
                f"{'true' if route.is_primary else 'false'} "
                "FROM catalog.event_classes AS ec "
                f"JOIN catalog.services AS s ON s.code = {sql_text(route.service_code)} "
                f"WHERE ec.event_number = {row.event_number} "
                "ON CONFLICT (event_class_id, service_id, source_column) DO UPDATE SET "
                "condition_code = EXCLUDED.condition_code, "
                "condition_label = EXCLUDED.condition_label, "
                "response_label = EXCLUDED.response_label, "
                "is_primary = EXCLUDED.is_primary;"
            )

    lines.append(
        "DELETE FROM catalog.classifier_version_events "
        "WHERE classifier_version_id = "
        f"(SELECT id FROM catalog.classifier_versions WHERE version_number = {version_number});"
    )
    lines.append(
        "INSERT INTO catalog.classifier_version_events "
        "(classifier_version_id, event_class_id) "
        "SELECT cv.id, ec.id FROM catalog.classifier_versions AS cv "
        "CROSS JOIN catalog.event_classes AS ec "
        f"WHERE cv.version_number = {version_number} "
        "AND ec.event_number IN ("
        + ",".join(str(row.event_number) for row in parsed.rows)
        + ") ON CONFLICT DO NOTHING;"
    )
    lines.append("COMMIT;")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


async def upsert_id(
    connection: AsyncConnection,
    table,
    values: dict,
    conflict_columns: list,
    update_values: dict,
) -> UUID:
    statement = pg_insert(table).values(id=uuid4(), **values)
    statement = statement.on_conflict_do_update(
        index_elements=conflict_columns,
        set_={**update_values, "updated_at": func.now()},
    ).returning(table.id)
    return (await connection.execute(statement)).scalar_one()


async def import_classifier(
    database_url: str,
    parsed: ParsedClassifier,
    workbook_path: Path,
    version_number: int,
) -> None:
    engine = create_async_engine(database_url)
    feature_1_labels, feature_2_labels, feature_3_labels = first_labels(parsed.rows)

    try:
        async with engine.begin() as connection:
            await connection.execute(
                update(ClassifierVersion)
                .where(ClassifierVersion.version_number != version_number)
                .values(is_active=False, updated_at=func.now())
            )
            version_id = await upsert_id(
                connection,
                ClassifierVersion,
                {
                    "version_number": version_number,
                    "name": "Классификатор происшествий v_046_11",
                    "source_name": workbook_path.name,
                    "description": (
                        "Импортированы A-D, G-I и маршруты служб N-CL. Строки с пометкой "
                        "'Не отображается оператору 112' исключены."
                    ),
                    "is_active": True,
                    "published_at": datetime.now(timezone.utc),
                },
                [ClassifierVersion.version_number],
                {
                    "name": "Классификатор происшествий v_046_11",
                    "source_name": workbook_path.name,
                    "description": (
                        "Импортированы A-D, G-I и маршруты служб N-CL. Строки с пометкой "
                        "'Не отображается оператору 112' исключены."
                    ),
                    "is_active": True,
                    "published_at": datetime.now(timezone.utc),
                },
            )
            await connection.execute(
                delete(classifier_version_events).where(
                    classifier_version_events.c.classifier_version_id == version_id
                )
            )

            service_ids: dict[str, UUID] = {}
            for code, name in sorted(parsed.service_names.items()):
                statement = pg_insert(Service).values(id=uuid4(), code=code, name=name)
                statement = statement.on_conflict_do_update(
                    index_elements=[Service.code], set_={"code": statement.excluded.code}
                ).returning(Service.id)
                service_ids[code] = (await connection.execute(statement)).scalar_one()

            event_type_ids: dict[int, UUID] = {}
            for code, name in sorted(parsed.group_names.items()):
                if code not in {row.event_type_code for row in parsed.rows}:
                    continue
                event_type_ids[code] = await upsert_id(
                    connection,
                    EventType,
                    {"code": code, "name": name},
                    [EventType.code],
                    {"name": name},
                )

            feature_1_ids: dict[tuple[int, int], UUID] = {}
            for key, label in feature_1_labels.items():
                type_code, code = key
                event_type_id = event_type_ids[type_code]
                feature_1_ids[key] = await upsert_id(
                    connection,
                    EventFeature1,
                    {
                        "event_type_id": event_type_id,
                        "code": code,
                        "statistics_name": label,
                        "operator_label": label,
                    },
                    [EventFeature1.event_type_id, EventFeature1.code],
                    {"statistics_name": label, "operator_label": label},
                )

            feature_2_ids: dict[tuple[int, int, int], UUID] = {}
            for key, label in feature_2_labels.items():
                parent_key = key[:2]
                feature_1_id = feature_1_ids[parent_key]
                feature_2_ids[key] = await upsert_id(
                    connection,
                    EventFeature2,
                    {
                        "event_feature_1_id": feature_1_id,
                        "code": key[2],
                        "name": label,
                    },
                    [EventFeature2.event_feature_1_id, EventFeature2.code],
                    {"name": label},
                )

            feature_3_ids: dict[tuple[int, int, int, int], UUID] = {}
            for key, label in feature_3_labels.items():
                parent_key = key[:3]
                feature_2_id = feature_2_ids[parent_key]
                feature_3_ids[key] = await upsert_id(
                    connection,
                    EventFeature3,
                    {
                        "event_feature_2_id": feature_2_id,
                        "code": key[3],
                        "name": label,
                    },
                    [EventFeature3.event_feature_2_id, EventFeature3.code],
                    {"name": label},
                )

            route_values: list[dict] = []
            service_links: set[tuple[UUID, UUID]] = set()
            for row in parsed.rows:
                labels = [
                    label
                    for label in (
                        row.feature_1_label,
                        row.feature_2_label,
                        row.feature_3_label,
                    )
                    if label
                ]
                event_name = " / ".join(labels)[:255]
                event_id = await upsert_id(
                    connection,
                    EventClass,
                    {
                        "event_number": row.event_number,
                        "event_type_id": event_type_ids[row.event_type_code],
                        "event_feature_1_id": feature_1_ids[row.key[:2]],
                        "event_feature_2_id": feature_2_ids[row.key[:3]],
                        "event_feature_3_id": feature_3_ids[row.key],
                        "feature_1_label": row.feature_1_label,
                        "feature_2_label": row.feature_2_label,
                        "feature_3_label": row.feature_3_label,
                        "name": event_name,
                        "main_service_id": service_ids.get(row.main_service_code),
                        "is_active": True,
                    },
                    [EventClass.event_number],
                    {
                        "feature_1_label": row.feature_1_label,
                        "feature_2_label": row.feature_2_label,
                        "feature_3_label": row.feature_3_label,
                        "name": event_name,
                        "main_service_id": service_ids.get(row.main_service_code),
                        "is_active": True,
                        "deleted_at": None,
                        "purge_after": None,
                    },
                )
                membership = pg_insert(classifier_version_events).values(
                    classifier_version_id=version_id,
                    event_class_id=event_id,
                )
                await connection.execute(
                    membership.on_conflict_do_nothing(
                        index_elements=[
                            classifier_version_events.c.classifier_version_id,
                            classifier_version_events.c.event_class_id,
                        ]
                    )
                )
                for route in row.service_routes:
                    service_id = service_ids[route.service_code]
                    service_links.add((event_id, service_id))
                    route_values.append(
                        {
                            "id": stable_id(
                                "service-route", row.event_number,
                                route.source_column, route.service_code,
                            ),
                            "event_class_id": event_id,
                            "service_id": service_id,
                            "source_column": route.source_column,
                            "condition_code": route.condition_code,
                            "condition_label": route.condition_label,
                            "response_label": route.response_label,
                            "is_primary": route.is_primary,
                        }
                    )

            links = [
                {"event_class_id": event_id, "service_id": service_id}
                for event_id, service_id in sorted(service_links)
            ]
            for offset in range(0, len(links), 500):
                statement = pg_insert(event_class_services).values(links[offset:offset + 500])
                await connection.execute(statement.on_conflict_do_nothing())
            for offset in range(0, len(route_values), 400):
                statement = pg_insert(EventServiceRoute).values(route_values[offset:offset + 400])
                await connection.execute(
                    statement.on_conflict_do_update(
                        index_elements=[
                            EventServiceRoute.event_class_id,
                            EventServiceRoute.service_id,
                            EventServiceRoute.source_column,
                        ],
                        set_={
                            "condition_code": statement.excluded.condition_code,
                            "condition_label": statement.excluded.condition_label,
                            "response_label": statement.excluded.response_label,
                            "is_primary": statement.excluded.is_primary,
                        },
                    )
                )
    finally:
        await engine.dispose()


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Импорт классификатора происшествий в PostgreSQL"
    )
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--version-number", type=int, default=46)
    action = parser.add_mutually_exclusive_group()
    action.add_argument(
        "--apply",
        action="store_true",
        help="Записать проверенные данные в БД; без флага выполняется только проверка",
    )
    action.add_argument(
        "--sql-output",
        type=Path,
        help="Сформировать повторно запускаемый SQL-файл для импорта",
    )
    return parser.parse_args()


def main() -> None:
    arguments = parse_arguments()
    workbook_path = arguments.workbook.resolve()
    parsed = parse_workbook(workbook_path)
    print(
        f"Проверено: {len(parsed.rows)} записей, "
        f"исключено: {parsed.excluded_rows}, "
        f"групп происшествий: {len(parsed.group_names)}, "
        f"служб: {len(parsed.service_names)}, "
        f"маршрутов: {sum(len(row.service_routes) for row in parsed.rows)}, "
        f"без главной службы в Excel: "
        f"{sum(row.main_service_code is None for row in parsed.rows)}"
    )

    if arguments.sql_output:
        output_path = arguments.sql_output.resolve()
        write_sql_import(
            output_path=output_path,
            parsed=parsed,
            workbook_path=workbook_path,
            version_number=arguments.version_number,
        )
        print(f"SQL для импорта создан: {output_path}")
        return

    if not arguments.apply:
        print("Проверка завершена. Для записи в БД добавьте --apply.")
        return

    database_url = os.getenv("DATABASE_URL")
    if not database_url:
        raise RuntimeError("Для импорта задайте переменную окружения DATABASE_URL")
    asyncio.run(
        import_classifier(
            database_url=database_url,
            parsed=parsed,
            workbook_path=workbook_path,
            version_number=arguments.version_number,
        )
    )
    print(f"Импортировано записей: {len(parsed.rows)}")


if __name__ == "__main__":
    main()
